"""Builds the bundled DEMONSTRATION templates as real .docx files with docxtpl placeholders.
Operators replace these with their own approved templates; the engine only needs a .docx whose
{{ field_id }} placeholders match a workspace config. Run `python -m app.templates.builder`."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

FILES_DIR = Path(__file__).parent / "files"
INK = RGBColor(0x1B, 0x1A, 0x17)
RULE = RGBColor(0x8A, 0x82, 0x72)

# Block DSL: ("title", t) ("h", t) ("p", t) ("fixed", t) ("kv", [(label, value)]) ("loop_li", list, fmt)
# ("loop_table", headers, list, cells) ("if", cond, [blocks]) ("note", t)
DECLARATION_TAIL = [
    ("p", "Jurisdiction: {{ jurisdiction }}"),
    ("p", "Date: {{ declaration_date }}"),
    ("note", "Signature of the declarant: to be added by the declarant in person. This draft carries no signature, "
             "witness, attestation, stamp or notarisation."),
]

TEMPLATES: dict[str, dict] = {
    "declaration_address": {
        "versions": ["1.0", "1.1"], "workspace": "legal", "title": "Address declaration (demonstration)",
        "jurisdiction": "Configurable", "banner": "DEMONSTRATION TEMPLATE. NOT LEGAL ADVICE. NOT LEGALLY VALID.",
        "blocks": [
            ("title", "DECLARATION OF ADDRESS"),
            ("fixed", "I, {{ applicant_name }}, state as follows:"),
            ("p", "1. My address, for the purpose of this declaration, is {{ applicant_address }}."),
            ("if", "previous_address", [("p", "2. My previous address was {{ previous_address }}.")]),
            ("fixed", "The statements in this declaration are made by me and are intended to be true to the best of my knowledge."),
            *DECLARATION_TAIL,
        ]},
    "declaration_name": {
        "versions": ["1.0"], "workspace": "legal", "title": "Name declaration (demonstration)",
        "jurisdiction": "Configurable", "banner": "DEMONSTRATION TEMPLATE. NOT LEGAL ADVICE. NOT LEGALLY VALID.",
        "blocks": [
            ("title", "DECLARATION OF NAME"),
            ("fixed", "I, {{ applicant_name }}, state as follows:"),
            ("p", "1. I was previously known as {{ previous_name }}."),
            ("fixed", "The statements in this declaration are made by me and are intended to be true to the best of my knowledge."),
            *DECLARATION_TAIL,
        ]},
    "declaration_general": {
        "versions": ["1.0"], "workspace": "legal", "title": "General declaration (demonstration)",
        "jurisdiction": "Configurable", "banner": "DEMONSTRATION TEMPLATE. NOT LEGAL ADVICE. NOT LEGALLY VALID.",
        "blocks": [
            ("title", "DECLARATION"),
            ("fixed", "I, {{ applicant_name }}, state as follows:"),
            ("p", "{{ declaration_statement }}"),
            ("p", "My address is {{ applicant_address }}."),
            ("fixed", "The statements in this declaration are made by me and are intended to be true to the best of my knowledge."),
            *DECLARATION_TAIL,
        ]},
    "legal_verification_checklist": {
        "versions": ["1.0"], "workspace": "legal", "title": "Operator verification checklist",
        "banner": "DEMONSTRATION TEMPLATE. OPERATOR WORKSHEET.",
        "blocks": [
            ("title", "OPERATOR VERIFICATION CHECKLIST"),
            ("kv", [("Applicant", "{{ applicant_name }}"), ("Declaration", "{{ declaration_purpose }}"),
                    ("Template", "{{ template_key }} v{{ template_version }}")]),
            ("h", "Before this draft is used"),
            ("p", "Identity document verified by operator: {{ identity_verified }}"),
            ("p", "Facts read back to the client: {{ facts_checked_with_client }}"),
            ("note", "Confirm the jurisdiction ({{ jurisdiction }}) and that the template version is the one your "
                     "business has approved."),
        ]},
    "rental_agreement": {
        "versions": ["1.0"], "workspace": "property", "title": "Rental agreement draft (demonstration)",
        "jurisdiction": "Configurable", "banner": "DEMONSTRATION DRAFT. NOT SIGNED OR EXECUTED. NOT LEGAL ADVICE.",
        "blocks": [
            ("title", "RENTAL AGREEMENT (DRAFT)"),
            ("fixed", "This draft sets out proposed terms between the landlord and the tenant named below."),
            ("kv", [("Landlord", "{{ landlord_name }}"), ("Landlord address", "{{ landlord_address }}"),
                    ("Tenant", "{{ tenant_name }}"), ("Tenant address", "{{ tenant_address }}"),
                    ("Property", "{{ property_address }}")]),
            ("h", "Terms"),
            ("p", "1. Tenancy begins on {{ start_date }} for {{ lease_months }} months, ending on {{ end_date }}."),
            ("p", "2. Monthly rent is {{ monthly_rent }}, payable by day {{ rent_due_day }} of each month."),
            ("p", "3. Security deposit is {{ security_deposit }}. Total payable at the start is {{ first_payment_total }}."),
            ("p", "4. Either party may end the tenancy by giving {{ notice_period_days }} days' written notice."),
            ("if", "furnished", [("p", "5. The property is let with the furnishings listed in the handover checklist.")]),
            ("fixed", "These terms are a draft for review by the parties and do not create obligations until agreed and signed by them."),
            ("note", "Signatures: to be added by the parties in person."),
        ]},
    "handover_checklist": {
        "versions": ["1.0"], "workspace": "property", "title": "Property handover checklist",
        "banner": "DEMONSTRATION TEMPLATE. OPERATOR WORKSHEET.",
        "blocks": [
            ("title", "PROPERTY HANDOVER CHECKLIST"),
            ("kv", [("Property", "{{ property_address }}"), ("Landlord", "{{ landlord_name }}"),
                    ("Tenant", "{{ tenant_name }}"), ("Handover date", "{{ start_date }}")]),
            ("h", "Items to inspect together"),
            ("loop_li", "inventory_items", "{{ item }}"),
            ("h", "Readings and keys"),
            ("p", "Electricity meter reading: ____________   Water meter reading: ____________"),
            ("p", "Keys handed over: ____________"),
        ]},
    "government_packet": {
        "versions": ["1.0"], "workspace": "government", "title": "Application preparation packet",
        "banner": "DEMONSTRATION TEMPLATE. NOT AN OFFICIAL FORM. NOT SUBMITTED.",
        "blocks": [
            ("title", "APPLICATION PREPARATION PACKET"),
            ("kv", [("Service", "{{ service_type }}"), ("Applicant", "{{ applicant_name }}"),
                    ("Date of birth", "{{ date_of_birth }}"), ("Phone", "{{ phone }}"), ("Address", "{{ applicant_address }}"),
                    ("Form used", "{{ form_name }}")]),
            ("h", "Answers mapped to the uploaded form"),
            ("loop_table", ["Form field", "Answer"], "form_mapping", ["{{ item.form_field }}", "{{ item.answer }}"]),
            ("h", "Supporting documents to collect"),
            ("loop_li", "checklist", "{{ item }}"),
            ("note", "Prepared for operator review. Nothing in this packet has been submitted to any authority."),
        ]},
    "government_intake": {
        "versions": ["1.0"], "workspace": "government", "title": "Customer intake sheet",
        "banner": "DEMONSTRATION TEMPLATE. INTERNAL INTAKE SHEET.",
        "blocks": [
            ("title", "CUSTOMER INTAKE SHEET"),
            ("kv", [("Service requested", "{{ service_type }}"), ("Applicant", "{{ applicant_name }}"),
                    ("Date of birth", "{{ date_of_birth }}"), ("Phone", "{{ phone }}"), ("Address", "{{ applicant_address }}")]),
            ("note", "Details as stated by the customer. Identity and eligibility are not checked by this system."),
        ]},
    "government_incomplete": {
        "versions": ["1.0"], "workspace": "government", "title": "Incomplete-field report",
        "banner": "DEMONSTRATION TEMPLATE. OPERATOR WORKSHEET.",
        "blocks": [
            ("title", "INCOMPLETE-FIELD REPORT"),
            ("kv", [("Applicant", "{{ applicant_name }}"), ("Service", "{{ service_type }}")]),
            ("h", "Form fields with no answer yet"),
            ("loop_li", "unmapped_form_fields", "{{ item }}"),
            ("h", "Supporting documents still to collect"),
            ("loop_li", "checklist", "{{ item }}"),
        ]},
    "travel_checklist": {
        "versions": ["1.0"], "workspace": "travel", "title": "Supporting-document checklist",
        "banner": "DEMONSTRATION TEMPLATE. NOT AN OFFICIAL REQUIREMENT LIST.",
        "blocks": [
            ("title", "SUPPORTING-DOCUMENT CHECKLIST"),
            ("kv", [("Applicant", "{{ applicant_name }}"), ("Destination", "{{ destination }}"),
                    ("Travel dates", "{{ trip_start }} to {{ trip_end }}"), ("Purpose", "{{ trip_purpose }}")]),
            ("h", "Documents to collect"),
            ("loop_li", "checklist", "{{ item }}"),
            ("h", "Source of these requirements"),
            ("p", "{{ requirements_source }}"),
            ("p", "Last verified: {{ requirements_verified }}"),
            ("note", "Always confirm current requirements with the destination's official source before the applicant relies on this list."),
        ]},
    "travel_itinerary": {
        "versions": ["1.0"], "workspace": "travel", "title": "Itinerary draft",
        "banner": "DEMONSTRATION TEMPLATE. DRAFT ONLY. NO BOOKINGS MADE.",
        "blocks": [
            ("title", "ITINERARY DRAFT"),
            ("kv", [("Traveller", "{{ applicant_name }}"), ("Destination", "{{ destination }}"),
                    ("Accommodation", "{{ accommodation }}"), ("Transport", "{{ transport_summary }}")]),
            ("loop_table", ["Day", "Date", "Plan"], "itinerary_days", ["{{ item.day }}", "{{ item.date }}", "{{ item.plan }}"]),
            ("note", "Only arrangements the applicant has confirmed belong here. This tool makes no reservations."),
        ]},
    "travel_cover_letter": {
        "versions": ["1.0"], "workspace": "travel", "title": "Travel cover-letter draft",
        "banner": "DEMONSTRATION TEMPLATE. DRAFT FOR THE APPLICANT TO REVIEW AND SIGN.",
        "blocks": [
            ("title", "COVER LETTER (DRAFT)"),
            ("p", "To the visa section,"),
            ("p", "I, {{ applicant_name }}, intend to travel to {{ destination }} from {{ trip_start }} to {{ trip_end }} "
                  "for the purpose of {{ trip_purpose }}."),
            ("p", "During my stay I will be at: {{ accommodation }}."),
            ("if", "occupation", [("p", "My occupation is {{ occupation }}.")]),
            ("p", "I am enclosing the documents listed in my application."),
            ("note", "Signature: to be added by the applicant in person. Review every statement before use."),
        ]},
    "travel_missing_info": {
        "versions": ["1.0"], "workspace": "travel", "title": "Missing-information report",
        "banner": "DEMONSTRATION TEMPLATE. OPERATOR WORKSHEET.",
        "blocks": [
            ("title", "MISSING-INFORMATION REPORT"),
            ("kv", [("Applicant", "{{ applicant_name }}"), ("Destination", "{{ destination }}")]),
            ("h", "Still needed from the applicant"),
            ("loop_li", "missing_items", "{{ item }}"),
        ]},
    "accounting_voucher": {
        "versions": ["1.0"], "workspace": "accounting", "title": "Expense voucher",
        "banner": "DEMONSTRATION TEMPLATE. INTERNAL VOUCHER. NOT AN ACCOUNTING-SYSTEM ENTRY.",
        "blocks": [
            ("title", "EXPENSE VOUCHER"),
            ("kv", [("Description", "{{ description }}"), ("Date", "{{ expense_date }}"), ("Supplier", "{{ supplier }}"),
                    ("Payment method", "{{ payment_method }}"), ("Category", "{{ category }}"),
                    ("Bill number", "{{ bill_number }}")]),
            ("h", "Amounts (calculated by code)"),
            ("kv", [("Base amount", "{{ base_amount }}"), ("Tax ({{ tax_rate }})", "{{ tax_amount }}"),
                    ("Total paid", "{{ total_amount }}")]),
            ("note", "Prepared for review. Nothing has been paid, filed or posted to an external system."),
        ]},
}


def template_file(key: str, version: str) -> Path:
    return FILES_DIR / f"{key}__{version}.docx"


def _run(par, text, bold=False, italic=False, size=11, color=INK, font="Cambria"):
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size, r.font.name, r.font.color.rgb = Pt(size), font, color
    return r


def _kv_table(doc, pairs):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in pairs:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(5), Cm(11)
        _run(cells[0].paragraphs[0], label, bold=True, size=10)
        _run(cells[1].paragraphs[0], value, size=10)
    doc.add_paragraph()


def _emit(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, b[1], bold=True, size=18)
        elif kind == "h":
            p = doc.add_paragraph(style="Heading 2")
            _run(p, b[1], bold=True, size=12)
        elif kind in ("p", "fixed"):
            _run(doc.add_paragraph(), b[1], size=11)
        elif kind == "note":
            _run(doc.add_paragraph(), b[1], italic=True, size=9.5, color=RULE)
        elif kind == "kv":
            _kv_table(doc, b[1])
        elif kind == "if":
            doc.add_paragraph("{%p if " + b[1] + " %}")
            _emit(doc, b[2])
            doc.add_paragraph("{%p endif %}")
        elif kind == "loop_li":
            doc.add_paragraph("{%p for item in " + b[1] + " %}")
            _run(doc.add_paragraph(style="List Bullet"), b[2], size=11)
            doc.add_paragraph("{%p endfor %}")
        elif kind == "loop_table":
            _, headers, listname, cells = b
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for c, h in zip(table.rows[0].cells, headers):
                _run(c.paragraphs[0], h, bold=True, size=10)
            table.add_row().cells[0].paragraphs[0].add_run("{%tr for item in " + listname + " %}")
            for c, t in zip(table.add_row().cells, cells):
                _run(c.paragraphs[0], t, size=10)
            table.add_row().cells[0].paragraphs[0].add_run("{%tr endfor %}")
            doc.add_paragraph()


def build_template(key: str, version: str, spec: dict) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    doc.styles["Normal"].font.name = "Cambria"
    _run(sec.header.paragraphs[0], spec["banner"], bold=True, size=8, color=RULE, font="Calibri")
    _run(sec.footer.paragraphs[0], f"Template {key} v{version}. Generated by Dastavez Desk. Draft pending human review.",
         size=8, color=RULE, font="Calibri")
    _emit(doc, spec["blocks"])
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    path = template_file(key, version)
    doc.save(path)
    return path


def ensure_templates(force: bool = False) -> None:
    for key, spec in TEMPLATES.items():
        for version in spec["versions"]:
            if force or not template_file(key, version).exists():
                build_template(key, version, spec)


if __name__ == "__main__":
    ensure_templates(force=True)
    print(f"Built {sum(len(s['versions']) for s in TEMPLATES.values())} templates in {FILES_DIR}")
