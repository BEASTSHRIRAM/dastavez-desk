"""Spreadsheet exports. Values come from deterministic code; nothing here calls a model."""
import csv
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LEDGER_HEADERS = ["Date", "Description", "Supplier", "Category", "Payment method", "Bill number",
                  "Base amount (INR)", "Tax rate (%)", "Tax (INR)", "Total paid (INR)", "Status"]


def _safe(v: Any) -> Any:
    """Spreadsheet formula-injection guard: dictated text must never be evaluated as a formula."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def ledger_row(values: dict[str, Any], computed: dict[str, Any], status: str) -> list[Any]:
    return [values.get("expense_date"), values.get("description"), values.get("supplier"), values.get("category"),
            values.get("payment_method"), values.get("bill_number") or "", computed.get("base_amount"),
            computed.get("tax_rate_pct"), computed.get("tax_amount"), computed.get("total_amount"), status]


def write_xlsx(path: Path, rows: list[list[Any]], headers: list[str] = LEDGER_HEADERS, title: str = "Ledger") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B1A17")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_safe(v) for v in row])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if isinstance(c.value, date):
                c.number_format = "dd-mmm-yyyy"
            elif isinstance(c.value, Decimal):
                c.number_format = "#,##0.00"
    for i, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1] if r[i - 1] is not None else "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 42)
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_csv(path: Path, rows: list[list[Any]], headers: list[str] = LEDGER_HEADERS) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel reads UTF-8
        w = csv.writer(fh)
        w.writerow(headers)
        for row in rows:
            w.writerow([v.isoformat() if isinstance(v, date) else (f"{v:.2f}" if isinstance(v, Decimal) else _safe(v))
                        for v in row])
    return path
