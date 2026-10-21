"""Template Studio: an advocate's own old draft becomes a reusable, approved template."""
import io

import pytest
from docx import Document

from app.services import studio
from tests.helpers import answer, approve, client, components, login, manifest

OLD_PARAS = [
    "AFFIDAVIT",
    "I, Mrs. Anita Verma, daughter of Shri Ramesh Verma, residing at 14 Lake View Road, Pune 411001, do hereby declare that "
    "I paid Rs. 18,000/- (Rupees Eighteen Thousand only) on 5th day of November, 2026 to Mr. Suresh Kulkar.",
    "Phone: 9876500011.",
    "The facts above are true to the best of my knowledge.",
    "Signature: ________",
]


def old_draft() -> bytes:
    d = Document()
    for t in OLD_PARAS:
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


DOCX = ("affidavit.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def analyze(h, data=None, name="affidavit.docx"):
    return client.post("/api/studio/analyze", files={"file": (name, data or old_draft(), DOCX[1])}, headers=h)


def pick(cands, skip=("Suresh Kulkar", "Mr.")):
    """Turn the analyst's proposals into selections, leaving some details out on purpose."""
    out = []
    for c in cands:
        if not c["include"] or c["text"] in skip or c["kind"] in ("blank",):
            continue
        out.append({k: c[k] for k in ("text", "field_id", "label", "type", "required", "high_impact", "occurrence", "link_to", "options")})
    return out


def build(h, up, selections, **kw):
    return client.post("/api/studio/build", json={"upload_id": up["upload_id"], "title": "Payment affidavit",
                                                  "selections": selections, **kw}, headers=h)


def test_analyst_proposes_the_variable_parts_of_an_old_draft():
    r = analyze(login())
    assert r.status_code == 200
    by_kind = {}
    for c in r.json()["candidates"]:
        by_kind.setdefault(c["kind"], []).append(c)
    assert {c["text"] for c in by_kind["name"]} == {"Anita Verma", "Ramesh Verma", "Suresh Kulkar"}
    assert {c["field_id"] for c in by_kind["name"]} >= {"declarant_name", "father_name"}
    assert [c["text"] for c in by_kind["address"]] == ["14 Lake View Road, Pune 411001"]
    assert [c["text"] for c in by_kind["amount"]] == ["Rs. 18,000/-"]
    assert by_kind["amount_words"][0]["link_to"] == by_kind["amount"][0]["field_id"]  # spelled out from its own amount
    assert by_kind["date"][0]["text"] == "5th day of November, 2026"
    assert [c["text"] for c in by_kind["phone"]] == ["9876500011"]
    sig = by_kind["blank"][0]
    assert sig["include"] is False and "signature" in sig["note"].lower()  # a signature line stays blank
    titles = {c["text"] for c in by_kind["title"]}
    assert {"Mrs.", "Shri", "Mr."} <= titles  # honorifics are choices, not baked in
    assert r.json()["agent"]["used"] is False  # demo mode: nothing was sent anywhere


@pytest.mark.parametrize("name,data,status", [
    ("draft.pdf", b"%PDF-1.4 fake", 422), ("notes.txt", b"hello", 422), ("broken.docx", b"PK\x03\x04 not really a docx", 422),
])
def test_studio_only_accepts_readable_word_files(name, data, status):
    assert analyze(login(), data, name).status_code == status


def test_old_draft_to_approved_template_to_a_new_clients_document_with_nothing_left_over():
    op, owner = login("operator@demo.test"), login("owner@demo.test")
    up = analyze(op).json()
    sels = pick(up["candidates"])
    r = build(op, up, sels, add_verification_checks=True)
    assert r.status_code == 201, r.text
    body = r.json()
    tpl = body["template"]
    assert tpl["approved"] is False and tpl["owned_by_business"] and tpl["version"] == "1.0"
    # the old client's remaining details are called out, never silently kept
    assert any("Suresh Kulkar" in w["text"] for w in body["leftovers"])
    text = " ".join(b.get("text", "") for b in body["preview"])
    assert "Anita Verma" not in text and "‹Declarant name›" in text  # the old name is gone from the template

    # nothing can start from it until the owner approves its wording
    blocked = client.post("/api/cases", json={"workspace": "custom", "template_key": tpl["key"]}, headers=op)
    assert blocked.status_code == 422
    assert client.post(f"/api/templates/{tpl['id']}/approve", headers=op).status_code == 403  # operators cannot sign off wording
    assert client.post(f"/api/templates/{tpl['id']}/approve", headers=owner).status_code == 200

    case = client.post("/api/cases", json={"workspace": "custom", "template_key": tpl["key"]}, headers=op).json()
    said = ("Prepare the affidavit. The declarant name is Neha Singh. The father name is Ravi Singh. "
            "The address is 9 Park Street, Mumbai 400001. The amount is 25000. The date is 12 November 2026. "
            "The phone is 9811122233.")
    assert client.post(f"/api/cases/{case['id']}/intake", json={"transcript": said}, headers=op).status_code == 202
    state = client.get(f"/api/cases/{case['id']}", headers=op).json()
    assert state["status"] == "awaiting_input"  # titles are never guessed, and the operator checks are always asked
    asked = [c["field_id"] for s in manifest(op, case["id"])["sections"] for c in s["components"]
             if "field_id" in c and not c.get("optional") and c["component"] not in ("source_evidence", "conflict_choice")]
    assert set(asked) == {"declarant_title", "father_title", "identity_verified", "facts_checked_with_client"}
    r = answer(op, case["id"], {"declarant_title": "Mrs.", "father_title": "Shri", "identity_verified": "yes",
                                "facts_checked_with_client": "yes"})
    assert r.json()["status"] == "review"

    # a title chosen wrongly is fixed in review, where it is visible and editable
    fixed = client.post(f"/api/cases/{case['id']}/review/edit", json={"answers": {"declarant_title": "Ms."}}, headers=op)
    assert fixed.status_code == 200
    approve(op, case["id"])
    out = client.post(f"/api/cases/{case['id']}/export", json={"formats": {"document": ["docx"]}}, headers=op).json()
    doc = Document(io.BytesIO(client.get(f"/api/documents/{out['documents'][0]['id']}/download", headers=op).content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "I, Ms. Neha Singh, daughter of Shri Ravi Singh, residing at 9 Park Street, Mumbai 400001" in text
    assert "Rs. 25,000.00 (Rupees Twenty Five Thousand only)" in text and "12 November 2026" in text
    assert "Anita Verma" not in text and "Ramesh Verma" not in text and "18,000" not in text  # no trace of the first client
    assert "Signature: ________" in text  # the signature line was left alone


def test_wording_that_states_a_figure_and_different_words_blocks_approval():
    """The advocate's own text says Rs. 18,000 but spells out eight thousand. The guard stops it."""
    d = Document()
    d.add_paragraph("Fee for Anita Verma is Rs. 18,000/- (Rupees Eight Thousand only), payable in full.")
    buf = io.BytesIO()
    d.save(buf)
    op, owner = login(), login("owner@demo.test")
    up = analyze(op, buf.getvalue()).json()
    guess = [c for c in up["candidates"] if c["text"] == "Anita Verma"]
    assert guess and guess[0]["include"] is False  # a plain name is only a suggestion, unticked
    name = [{k: guess[0][k] for k in ("text", "field_id", "label", "type", "required", "high_impact", "occurrence", "link_to", "options")}]
    tpl = build(op, up, name).json()["template"]
    client.post(f"/api/templates/{tpl['id']}/approve", headers=owner)
    case = client.post("/api/cases", json={"workspace": "custom", "template_key": tpl["key"]}, headers=op).json()
    client.post(f"/api/cases/{case['id']}/intake", json={"transcript": "The person name is Neha Singh."}, headers=op)
    assert client.get(f"/api/cases/{case['id']}", headers=op).json()["status"] == "review"
    panel = components(manifest(op, case["id"]), "approval_panel")[0]
    assert any(i["code"] == "amount_mismatch" and i["severity"] == "error" for i in panel["issues"])
    assert panel["can_approve"] is False
    refused = client.post(f"/api/cases/{case['id']}/approve", json={"confirm": [f["field_id"] for f in panel["confirm_fields"]]}, headers=op)
    assert refused.status_code == 409


def test_template_building_rejects_unsafe_or_inconsistent_choices():
    op = login()
    up = analyze(op).json()
    good = pick(up["candidates"])
    amount = next(s for s in good if s["type"] == "currency")
    bad_ids = [("Capital", "must start"), ("1abc", "must start"), ("checklist", "reserved"), ("x", "must start"), ("has space", "must start"),
               ("name_words", "reserved"), ("done_flag", "reserved")]
    for fid, _ in bad_ids:
        sels = [{**amount, "field_id": fid}]
        assert build(op, up, sels).status_code == 422, fid
    clash = [{**amount, "field_id": "same"}, {**good[0], "field_id": "same"}]
    assert build(op, up, clash).status_code == 422  # one id cannot be both a name and a currency
    assert build(op, up, []).status_code == 422  # nothing varies, so it is not a template
    assert client.post("/api/studio/build", json={"upload_id": "0" * 32, "title": "Nope", "selections": good}, headers=op).status_code == 404
    words = next(s for s in good if s["type"] == "amount_words")
    assert build(op, up, [words]).status_code == 422  # the words need their amount selected too


def test_templates_belong_to_one_business_and_viewers_cannot_use_the_studio():
    op, owner, other, viewer = login(), login("owner@demo.test"), login("other@demo.test"), login("viewer@demo.test")
    up = analyze(op).json()
    tpl = build(op, up, pick(up["candidates"])).json()["template"]
    assert client.get(f"/api/templates/{tpl['id']}", headers=other).status_code == 404
    assert client.get(f"/api/templates/{tpl['id']}/preview", headers=other).status_code == 404
    assert client.post(f"/api/templates/{tpl['id']}/approve", headers=login("other@demo.test")).status_code in (403, 404)
    assert tpl["id"] not in [t["id"] for t in client.get("/api/templates", headers=other).json()]
    assert client.post("/api/studio/build", json={"upload_id": up["upload_id"], "title": "Steal", "selections": pick(up["candidates"])},
                       headers=other).status_code == 404  # another business cannot build from this upload
    assert analyze(viewer).status_code == 403
    assert client.post("/api/cases", json={"workspace": "custom", "template_key": tpl["key"]}, headers=other).status_code == 422
    client.post(f"/api/templates/{tpl['id']}/approve", headers=owner)
    assert client.post("/api/cases", json={"workspace": "custom", "template_key": tpl["key"]}, headers=other).status_code == 422
    # bundled templates cannot be edited by anyone
    builtin = next(t for t in client.get("/api/templates", headers=owner).json() if not t["owned_by_business"])
    assert client.post(f"/api/templates/{builtin['id']}/approve", headers=owner).status_code == 403


def test_a_new_version_does_not_replace_the_approved_one_until_it_is_approved():
    op, owner = login(), login("owner@demo.test")
    up = analyze(op).json()
    v1 = build(op, up, pick(up["candidates"])).json()["template"]
    client.post(f"/api/templates/{v1['id']}/approve", headers=owner)
    v2 = build(op, up, pick(up["candidates"]), template_key=v1["key"]).json()["template"]
    assert v1["version"] == "1.0" and v2["version"] == "1.1" and v2["key"] == v1["key"] and v2["approved"] is False
    case = client.post("/api/cases", json={"workspace": "custom", "template_key": v1["key"]}, headers=op).json()
    client.post(f"/api/cases/{case['id']}/intake", json={"transcript": "The declarant name is Neha Singh."}, headers=op)
    pv = client.get(f"/api/cases/{case['id']}/manifest", headers=op).json()
    assert pv["stage"] == "questions"  # the case runs on the approved 1.0, not the unapproved 1.1
    client.post(f"/api/templates/{v1['id']}/retire", headers=owner)
    assert client.post("/api/cases", json={"workspace": "custom", "template_key": v1["key"]}, headers=op).status_code == 422


# ------------------------------------------------------------------ the optional analyst model
class FakeAnalyst:
    name = "fake"

    def __init__(self, payload):
        self.payload = payload

    def complete_json(self, system, user):
        return self.payload


def test_the_analyst_model_may_relabel_and_add_spans_but_only_ones_that_exist_in_the_document():
    import tempfile
    from pathlib import Path
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "d.docx"
        p.write_bytes(old_draft())
        cands, _ = studio.analyze(p)
        text = studio.full_text(Document(str(p)))
    fake = FakeAnalyst({"labels": [{"text": "Anita Verma", "label": "Deponent full name", "type": "person"}],
                        "extra": [{"text": "Lake View Road", "label": "Street", "type": "text"},   # inside an existing candidate: dropped
                                  {"text": "true to the best of my knowledge", "label": "Closing", "type": "text"},
                                  {"text": "Not in the document at all", "label": "Ghost", "type": "text"}]})
    out, notes = studio.agent_refine(cands, text, fake)
    assert next(c for c in out if c.text == "Anita Verma").label == "Deponent full name"
    added = [c for c in out if c.source == "agent" and c.kind == "agent"]
    assert [c.text for c in added] == ["true to the best of my knowledge"]
    assert sum("word-for-word" in n for n in notes) == 2
