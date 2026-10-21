"""One end-to-end test per workspace, all through the same engine."""
import io

from openpyxl import load_workbook

from tests.helpers import answer, approve, asked_fields, client, components, login, manifest, new_case


def download(h, doc_id):
    r = client.get(f"/api/documents/{doc_id}/download", headers=h)
    assert r.status_code == 200, r.text
    return r.content


def export(h, case_id, formats):
    r = client.post(f"/api/cases/{case_id}/export", json={"formats": formats}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def test_legal_affidavit_asks_only_what_is_missing_and_preserves_approved_wording():
    h = login()
    text = ("Prepare a declaration using our approved template. The applicant is Anita Verma, the new address is "
            "14 Lake View Road, Pune. Change the address and date, but don't change the approved wording. "
            "Ask me if anything required is missing.")
    case = new_case(h, "legal", text)
    assert case["status"] == "awaiting_input"
    m = manifest(h, case["id"])
    asked = asked_fields(m)
    assert {"declaration_date", "jurisdiction", "identity_verified", "facts_checked_with_client"} <= set(asked)
    assert "applicant_name" not in asked and "applicant_address" not in asked
    # The declaration type was heard ("change the address") but is only a suggestion; it is not asked unless uncertain.
    r = answer(h, case["id"], {"declaration_date": "2026-11-05", "jurisdiction": "Maharashtra",
                               "identity_verified": "yes", "facts_checked_with_client": "yes"})
    assert r.status_code == 200 and r.json()["status"] == "review"
    m = manifest(h, case["id"])
    comparison = components(m, "change_comparison")[0]["sections"]
    fixed = [s for s in comparison if not s["changed"]]
    assert any("true to the best of my knowledge" in s["approved"] for s in fixed)
    assert all(s["approved"] == s["draft"] for s in fixed)  # approved wording is untouched
    assert any(s["changed"] and "Anita Verma" in s["draft"] for s in comparison)
    approve(h, case["id"])
    out = export(h, case["id"], {"declaration": ["docx"]})
    assert [d["fmt"] for d in out["documents"]] == ["docx"]
    assert download(h, out["documents"][0]["id"]).startswith(b"PK")


def test_rental_agreement_and_handover_checklist_from_same_case_data():
    h = login()
    case = new_case(h, "property")
    assert case["status"] == "awaiting_input"
    m = manifest(h, case["id"])
    fields = {c["field_id"]: c for c in components(m) if "field_id" in c and c["component"] != "source_evidence"}
    assert "monthly_rent" not in fields  # already extracted, not re-asked
    options = fields["landlord_party"]["party_options"]
    assert {o["label"] for o in options} == {"Meera Kapoor", "Arjun Desai"}  # only this business's records
    meera = next(o["id"] for o in options if o["label"] == "Meera Kapoor")
    r = answer(h, case["id"], {"landlord_party": meera, "property_address": "Flat 4B, Green Heights, Pune",
                               "lease_months": 11, "furnished": "yes"})
    assert r.status_code == 200
    # Furnished was answered yes, so the conditional question appears (and only now).
    assert r.json()["status"] == "awaiting_input"
    assert "inventory_summary" in asked_fields(manifest(h, case["id"]))
    r = answer(h, case["id"], {"inventory_summary": "Bed\nWardrobe\nSofa"})
    assert r.json()["status"] == "review"
    approve(h, case["id"])
    out = export(h, case["id"], {"rental_agreement": ["docx"], "handover_checklist": ["docx"]})
    assert {d["doc_key"] for d in out["documents"]} == {"rental_agreement", "handover_checklist"}
    from docx import Document
    doc = Document(io.BytesIO(download(h, next(d["id"] for d in out["documents"] if d["doc_key"] == "rental_agreement"))))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r_ in t.rows for c in r_.cells)
    assert "Rs. 18,000.00" in text and "Rs. 50,000.00" in text and "Rs. 68,000.00" in text
    assert "1 November 2026" in text and "30 September 2027" in text and "Rahul Sharma" in text
    assert "Meera Kapoor" in text and "Arjun Desai" not in text
    checklist = Document(io.BytesIO(download(h, next(d["id"] for d in out["documents"] if d["doc_key"] == "handover_checklist"))))
    assert "Sofa" in "\n".join(p.text for p in checklist.paragraphs)


def test_government_application_maps_uploaded_form_and_flags_missing_documents():
    h = login()
    case = new_case(h, "government")
    assert case["status"] == "review"  # everything required was dictated
    r = client.post(f"/api/cases/{case['id']}/attachments/sample-form", headers=h)
    assert r.status_code == 201
    form = r.json()
    assert "Full name" in form["fields"] and form["mapping"]["Full name"] == "applicant_name"
    assert "Father's name" not in form["mapping"]  # unknown stays unmapped for the operator
    m = manifest(h, case["id"])
    assert any("Proof of new address" in i for c in components(m, "document_checklist") for i in c["items"])
    approve(h, case["id"])
    out = export(h, case["id"], {"application_packet": ["docx"], "incomplete_report": ["docx"]})
    from docx import Document
    report = Document(io.BytesIO(download(h, next(d["id"] for d in out["documents"] if d["doc_key"] == "incomplete_report"))))
    assert "Father's name" in "\n".join(p.text for p in report.paragraphs)


def test_travel_checklist_shows_source_and_never_invents_requirements():
    h = login()
    case = new_case(h, "travel")
    assert case["status"] == "awaiting_input"
    m = manifest(h, case["id"])
    assert {"nationality"} <= set(asked_fields(m))
    r = answer(h, case["id"], {"nationality": "Indian", "transport_summary": "Flight booked by applicant"})
    assert r.json()["status"] == "review"
    m = manifest(h, case["id"])
    checklist = components(m, "document_checklist")[0]
    assert "Demonstration dataset" in checklist["source"]
    assert checklist["last_verified"] == "Not verified"
    itinerary = next(c for c in components(m, "document_preview") if c["doc_key"] == "itinerary_draft")
    rows = [r_ for b in itinerary["blocks"] if b["type"] == "table" for r_ in b["rows"]]
    assert len([r_ for r_ in rows if r_[0].isdigit()]) == 10  # 1-10 December inclusive


def test_unknown_destination_gets_a_warning_and_no_checklist():
    h = login()
    text = ("Prepare documents for applicant Ravi Iyer travelling to Atlantis from 1 December 2026 until 5 December 2026 "
            "for tourism, staying at Coral Inn. Nationality is Indian.")
    case = new_case(h, "travel", text)
    m = manifest(h, case["id"])
    panel = components(m, "approval_panel")[0]
    assert any("No administrator-approved requirements" in i["message"] for i in panel["issues"])
    assert not components(m, "document_checklist")


def test_bookkeeping_expense_uses_exact_decimals_and_exports_valid_spreadsheet():
    h = login()
    case = new_case(h, "accounting")
    assert case["status"] == "awaiting_input"
    m = manifest(h, case["id"])
    assert asked_fields(m) == ["supplier"]  # amount, date, method, category all heard
    r = answer(h, case["id"], {"supplier": "Sharma Office Supplies", "tax_rate": "18%", "amount_includes_tax": "yes"})
    assert r.status_code == 200
    m = manifest(h, case["id"])
    assert components(m, "approval_panel")[0]["can_approve"]
    approve(h, case["id"])
    out = export(h, case["id"], {"ledger_entry": ["xlsx", "csv"]})
    xlsx = next(d for d in out["documents"] if d["fmt"] == "xlsx")
    wb = load_workbook(io.BytesIO(download(h, xlsx["id"])))
    row = [c.value for c in wb.active[2]]
    from decimal import Decimal
    assert Decimal(str(row[6])) == Decimal("3559.32") and Decimal(str(row[8])) == Decimal("640.68")
    assert Decimal(str(row[9])) == Decimal("4200.00")  # 4200 inclusive of 18%: base + tax == total exactly
    assert wb.active["J2"].number_format == "#,##0.00"
    csv_doc = next(d for d in out["documents"] if d["fmt"] == "csv")
    assert b"3559.32" in download(h, csv_doc["id"])


def test_missing_details_come_with_a_ready_to_send_client_message():
    h = login()
    case = new_case(h, "legal", "Prepare a declaration. The applicant is Anita Verma, the new address is 14 Lake View Road, Pune.")
    m = manifest(h, case["id"])
    [note] = components(m, "client_message")
    assert "Declaration date" in note["text"] and "Jurisdiction" in note["text"]
    assert "Operator" not in note["text"] and "read these facts back" not in note["text"]  # operator-only checks stay out
    gov = new_case(h, "government")
    client.post(f"/api/cases/{gov['id']}/attachments/sample-form", headers=h)
    review = manifest(h, gov["id"])
    [msg] = components(review, "client_message")
    assert "Please send or bring" in msg["text"] and "Proof of new address" in msg["text"]
